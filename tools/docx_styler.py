import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

NAVY = RGBColor(30, 58, 138)       # #1E3A8A
SLATE = RGBColor(15, 23, 42)      # #0F172A
TEAL = RGBColor(13, 148, 136)     # #0D9488
DARK_GRAY = RGBColor(51, 65, 85)  # #334155
MUTED = RGBColor(100, 116, 139)   # #64748B
RED = RGBColor(225, 29, 72)       # #E11D48


def set_document_margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("ReguGraph-UAE | Autonomous CBUAE Regulatory Intelligence Platform")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = MUTED

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Confidential & Proprietary — Prepared for Official Demonstration on October 19, 2026")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED


def add_title(doc, title, subtitle, metadata_dict=None):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(title)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = SLATE

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = TEAL

    if metadata_dict:
        table = doc.add_table(rows=len(metadata_dict), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        r_idx = 0
        for k, v in metadata_dict.items():
            cell_k = table.cell(r_idx, 0)
            cell_v = table.cell(r_idx, 1)
            cell_k.width = Inches(2.2)
            cell_v.width = Inches(4.3)
            
            pk = cell_k.paragraphs[0]
            pk.paragraph_format.space_after = Pt(2)
            rk = pk.add_run(k)
            rk.font.name = "Calibri"
            rk.font.size = Pt(9.5)
            rk.font.bold = True
            rk.font.color.rgb = SLATE

            pv = cell_v.paragraphs[0]
            pv.paragraph_format.space_after = Pt(2)
            rv = pv.add_run(str(v))
            rv.font.name = "Calibri"
            rv.font.size = Pt(9.5)
            rv.font.color.rgb = DARK_GRAY

            # Shading
            shading_elm = parse_xml(r'<w:shd {} w:fill="F8FAFC"/>'.format(nsdecls('w')))
            cell_k._tc.get_or_add_tcPr().append(shading_elm)
            shading_elm2 = parse_xml(r'<w:shd {} w:fill="FFFFFF"/>'.format(nsdecls('w')))
            cell_v._tc.get_or_add_tcPr().append(shading_elm2)
            r_idx += 1

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(14)


def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = NAVY
    return p


def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = TEAL
    return p


def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = SLATE
    return p


def add_paragraph(doc, text, bold_prefix=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = SLATE
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = DARK_GRAY
    r_text.font.italic = italic
    return p


def add_bullet(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = SLATE
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = DARK_GRAY
    return p


def add_callout(doc, text, title="KEY REGULATORY PRINCIPLE"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)

    # Shading and left border
    shading_elm = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading_elm)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        r'<w:tcBorders {} >'
        r'<w:top w:val="none"/>'
        r'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        r'<w:bottom w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'</w:tcBorders>'.format(nsdecls('w'))
    )
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10.0)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10.0)
    r_text.font.color.rgb = DARK_GRAY

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(6)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        shd = parse_xml(r'<w:shd {} w:fill="1E3A8A"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)

    # Data Rows
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.0)
                r.font.color.rgb = DARK_GRAY
            shd = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg_color))
            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

    # Set Column Widths if provided
    if col_widths:
        for r in table.rows:
            for idx, w in enumerate(col_widths):
                if idx < len(r.cells):
                    r.cells[idx].width = Inches(w)

    # Apply borders
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        r'<w:tblBorders {} >'
        r'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        r'<w:left w:val="none"/>'
        r'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        r'<w:right w:val="none"/>'
        r'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        r'<w:insideV w:val="none"/>'
        r'</w:tblBorders>'.format(nsdecls('w'))
    )
    tblPr.append(borders)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(6)
    return table


def add_image_with_caption(doc, image_path, caption, width_inches=6.3):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(f"Figure: {caption}")
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9.0)
        run_cap.font.italic = True
        run_cap.font.color.rgb = MUTED
